"""Testes de autenticacao/autorizacao e da leitura da planilha SESA - rodam sem banco.

Uso (na raiz do projeto): python3 -m unittest discover -s tests -v
"""
import io
import os
import sys
import unittest
from unittest import mock

# app sem banco: chave de sessao fixa e sem a migracao de inicio
os.environ.setdefault("APP_SECRET_KEY", "chave-de-teste")
os.environ["APP_SEM_MIGRACAO_INICIO"] = "1"
os.environ.pop("USUARIO_TESTE", None)
os.environ.pop("SENHA_TESTE", None)
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import app as modulo  # noqa: E402
from werkzeug.security import generate_password_hash  # noqa: E402

SENHA = "senha-certa"


def _sistemausuario(administrador=False, senha=SENHA, perfil=1):
    return {
        "sistemausuario": 1,
        "senha_hash": generate_password_hash(senha) if senha else None,
        "tela_padrao": None,
        "sistemaperfil": perfil,
        "perfil_nome": "Administrador" if administrador else "Supervisor",
        "administrador": administrador,
    }


class BaseTeste(unittest.TestCase):
    def setUp(self):
        modulo.app.config["TESTING"] = True
        self.cliente = modulo.app.test_client()
        modulo._tentativas_login.clear()
        self.patches = [
            mock.patch.object(modulo, "_erro_conexao_banco", return_value=None),
            mock.patch.object(modulo, "_paginas_do_perfil", return_value=["estoque"]),
        ]
        for patch in self.patches:
            patch.start()

    def tearDown(self):
        for patch in self.patches:
            patch.stop()

    def logar(self, identificacao="joao", administrador=False):
        with mock.patch.object(modulo, "run_query", return_value=[{"identificacao": identificacao, "nome": "Joao"}]), \
                mock.patch.object(modulo, "_buscar_sistemausuario", return_value=_sistemausuario(administrador)):
            resposta = self.cliente.post("/api/login", json={"usuario": identificacao, "senha": SENHA})
        self.assertEqual(resposta.status_code, 200, resposta.get_json())
        return resposta


class TestSessao(BaseTeste):
    def test_api_sem_sessao_recusa(self):
        self.assertEqual(self.cliente.get("/api/armazens").status_code, 401)
        self.assertEqual(self.cliente.post("/api/perfis", json={}).status_code, 401)
        self.assertEqual(self.cliente.put("/api/configuracao_banco", json={}).status_code, 401)

    def test_paginas_html_continuam_publicas(self):
        self.assertEqual(self.cliente.get("/").status_code, 200)

    def test_login_cria_sessao_e_logout_encerra(self):
        self.logar()
        with mock.patch.object(modulo, "run_query", return_value=[]):
            self.assertEqual(self.cliente.get("/api/armazens").status_code, 200)
        self.cliente.post("/api/logout")
        self.assertEqual(self.cliente.get("/api/armazens").status_code, 401)

    def test_rota_de_admin_exige_admin(self):
        self.logar(administrador=False)
        self.assertEqual(self.cliente.get("/api/paginas_sistema").status_code, 403)
        self.assertEqual(self.cliente.post("/api/usuarios_sistema/resetar_senha", json={"identificacao": "x"}).status_code, 403)

    def test_rota_de_admin_libera_admin(self):
        self.logar(administrador=True)
        self.assertEqual(self.cliente.get("/api/paginas_sistema").status_code, 200)

    def test_usuario_logado_do_navegador_eh_ignorado(self):
        # quem decide a permissao eh a sessao, nao o parametro usuario_logado
        self.logar(identificacao="joao")
        with mock.patch.object(modulo, "_usuario_pode_editar_processo", return_value=False) as checagem:
            self.cliente.get("/api/processos_hospital/permissoes?usuario_logado=admin")
        checagem.assert_called_once_with("joao")


class TestLogin(BaseTeste):
    def test_usuario_inexistente_e_senha_errada_mesma_resposta(self):
        with mock.patch.object(modulo, "run_query", return_value=[]):
            inexistente = self.cliente.post("/api/login", json={"usuario": "ninguem", "senha": "x"})
        with mock.patch.object(modulo, "run_query", return_value=[{"identificacao": "joao", "nome": "Joao"}]), \
                mock.patch.object(modulo, "_buscar_sistemausuario", return_value=_sistemausuario()):
            senha_errada = self.cliente.post("/api/login", json={"usuario": "joao", "senha": "errada"})
        self.assertEqual(inexistente.status_code, 401)
        self.assertEqual(senha_errada.status_code, 401)
        self.assertEqual(inexistente.get_json()["msg"], senha_errada.get_json()["msg"])

    def test_bloqueia_apos_tentativas_erradas(self):
        with mock.patch.object(modulo, "run_query", return_value=[]):
            for _ in range(modulo.LIMITE_TENTATIVAS_LOGIN):
                self.cliente.post("/api/login", json={"usuario": "alvo", "senha": "x"})
            bloqueado = self.cliente.post("/api/login", json={"usuario": "alvo", "senha": "x"})
        self.assertEqual(bloqueado.status_code, 429)

    def test_sem_usuario_de_teste_configurado_banco_fora_nao_entra(self):
        with mock.patch.object(modulo, "_erro_conexao_banco", return_value="sem conexao"):
            resposta = self.cliente.post("/api/login", json={"usuario": "teste", "senha": "teste"})
        self.assertEqual(resposta.status_code, 503)

    def test_usuario_de_teste_so_acessa_configuracao(self):
        with mock.patch.object(modulo, "USUARIO_TESTE", "suporte"), mock.patch.object(modulo, "SENHA_TESTE", "s3nh4"), \
                mock.patch.object(modulo, "_erro_conexao_banco", return_value="sem conexao"):
            resposta = self.cliente.post("/api/login", json={"usuario": "suporte", "senha": "s3nh4"})
            self.assertEqual(resposta.status_code, 200)
            self.assertEqual(self.cliente.get("/api/armazens").status_code, 403)
            self.assertEqual(self.cliente.get("/api/configuracao_banco").status_code, 200)


class TestDefinirSenha(BaseTeste):
    def test_nao_troca_senha_de_quem_ja_tem(self):
        # UPDATE condicionado a senha_hash IS NULL nao atualiza nada -> recusa
        with mock.patch.object(modulo, "run_query", return_value=[]) as consulta:
            resposta = self.cliente.post("/api/definir_senha", json={"identificacao": "admin", "senha": "nova-senha"})
        self.assertEqual(resposta.status_code, 403)
        self.assertIn("senha_hash IS NULL", consulta.call_args[0][0])

    def test_primeiro_acesso_define(self):
        with mock.patch.object(modulo, "run_query", return_value=[{"sistemausuario": 1}]):
            resposta = self.cliente.post("/api/definir_senha", json={"identificacao": "novo", "senha": "nova-senha"})
        self.assertEqual(resposta.status_code, 200)


class TestConfiguracaoBanco(BaseTeste):
    def test_trocar_host_exige_senha(self):
        self.logar(administrador=True)
        with mock.patch.dict(modulo.DB_CONFIG, {"host": "db-real", "port": 5432, "user": "postgres", "password": "x"}), \
                mock.patch.object(modulo, "salvar_config") as salvar:
            resposta = self.cliente.put("/api/configuracao_banco", json={
                "host": "servidor-de-fora", "port": 5432, "dbname": "x", "user": "postgres", "password": ""})
        self.assertEqual(resposta.status_code, 400)
        salvar.assert_not_called()


class TestLimites(BaseTeste):
    def test_upload_acima_do_limite_recusado(self):
        self.logar()
        grande = io.BytesIO(b"x" * (modulo.app.config["MAX_CONTENT_LENGTH"] + 1))
        resposta = self.cliente.post("/api/processos_sesa/importar", data={"arquivo": (grande, "a.xlsx")},
                                     content_type="multipart/form-data")
        self.assertEqual(resposta.status_code, 413)


class TestPagamentos(BaseTeste):
    def setUp(self):
        super().setUp()
        for patch in (mock.patch.object(modulo, "_garantir_tabelas_pagamento"),
                      mock.patch.object(modulo, "_destaque_existe", return_value=True)):
            patch.start()
            self.patches.append(patch)
        self.logar()

    def test_prioridade_fora_de_1_a_3_recusada(self):
        with mock.patch.object(modulo, "run_query") as consulta:
            resposta = self.cliente.put("/api/pagamentos/10", json={"prioridade": "4"})
        self.assertEqual(resposta.status_code, 400)
        consulta.assert_not_called()

    def test_data_de_envio_invalida_recusada(self):
        with mock.patch.object(modulo, "run_query"):
            resposta = self.cliente.put("/api/pagamentos/10", json={"data_envio_empenho": "31/02/2026"})
        self.assertEqual(resposta.status_code, 400)

    def test_pago_exige_numero_de_empenho(self):
        with mock.patch.object(modulo, "run_query", return_value=[{"numero_empenho": None}]):
            resposta = self.cliente.post("/api/pagamentos/10/situacao", json={"acao": "pago"})
        self.assertEqual(resposta.status_code, 400)

    def test_edicao_grava_usuario_da_sessao(self):
        with mock.patch.object(modulo, "run_query", return_value=[]) as consulta:
            resposta = self.cliente.put("/api/pagamentos/10", json={"numero_empenho": "NE1", "prioridade": "2"})
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(consulta.call_args[0][1][-1], "joao")


@unittest.skipUnless(modulo.OPENPYXL_DISPONIVEL, "openpyxl nao instalado")
class TestPlanilhaSesa(unittest.TestCase):
    def _planilha(self):
        from openpyxl import Workbook
        pasta = Workbook()
        pasta.active.title = "LEG-STATUS"
        aba = pasta.create_sheet("MED. 2024-2026")
        aba.append(["ANO", "PROCESSO", "PLAN", "PLAN", "Nº ITEM", "CODIGO SRP", "DESCRIÇÃO", "STATUS"])
        aba.append([2024, "24001.1/2024", "2024/1", "J", 1, 1563306, "AFLIBERCEPTE", "FINALIZADO"])
        aba.append([None, None, None, None, None, None, None, None])
        outra = pasta.create_sheet("CONTRATO")
        outra.append(["ANO"])
        buffer = io.BytesIO()
        pasta.save(buffer)
        return buffer.getvalue()

    def test_grupos_colunas_e_abas_ignoradas(self):
        linhas, resumo = modulo._ler_planilha_sesa(self._planilha())
        self.assertEqual(len(linhas), 1)
        item = linhas[0]
        self.assertEqual(item["sigla"], "MED")
        self.assertEqual(item["codigo_srp"], "1563306")
        self.assertEqual(item["plan"], "2024/1")
        self.assertEqual(item["curva"], "J")  # 2o "PLAN" da aba MED eh a curva
        self.assertIn("LEG-STATUS", resumo["abas_ignoradas"])
        self.assertIn("CONTRATO", resumo["abas_ignoradas"])

    def test_ordem_natural_do_numero_do_item(self):
        itens = [{"numero_item": n} for n in ["10", "2", None, "G1 - 10", "G1 - 2", "1"]]
        ordem = [i["numero_item"] for i in sorted(itens, key=modulo._chave_numero_item)]
        self.assertEqual(ordem, ["1", "2", "10", "G1 - 2", "G1 - 10", None])


if __name__ == "__main__":
    unittest.main()
